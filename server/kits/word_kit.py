#!/usr/bin/env python3
"""
word-kit.py — Mesin Pembuat Dokumen Word (.docx) Pro dari JSON Spec (Python Engine)
Kompatibel 100% dengan spec.json milik docx-kit.js + fitur mendalam:
  - Smart Image Payload & Placement:
      placement: "center" | "full-banner" | "left" | "right" | "split-left" | "split-right"
      side_text / side_bullets untuk gambar berdampingan dengan teks secara rapi!
      aspect_ratio ("16:9", "4:3", "1:1") + fit_mode ("cover" / "contain")
  - Blok "gallery": Grid gambar 2 atau 3 kolom sejajar proporsional
  - Blok "curve" / "chart": Langsung merender grafik kurva 300-DPI via viz_kit.py ke dalam dokumen
  - Blok "kpi_row": Kartu metrik KPI 2-4 kolom sejajar
  - Blok lengkap: title, h1, h2, h3, p, bullets, numbers, table, kv, callout, quote, code,
                  image, gallery, curve, kpi_row, pagebreak, spacer, rule, sectionhead, lr, letterhead, signature

Cara pakai (CLI):
  python3 word-kit.py spec.json output.docx
"""

import json
import os
import re
import sys
import tempfile
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import viz_kit

THEMES = {
    "navy":   {"primary": "1F3A5F", "accent": "2E75B6", "callout": "EEF4FB", "zebra": "F4F7FA", "coverSub": "C9DAF0"},
    "green":  {"primary": "1B5E4A", "accent": "3E9B7A", "callout": "EDF7F3", "zebra": "F3F9F6", "coverSub": "CDEBDD"},
    "maroon": {"primary": "7A1F2B", "accent": "B5485D", "callout": "FBEFF1", "zebra": "FAF3F4", "coverSub": "F0CDD3"},
    "tech":   {"primary": "22223B", "accent": "4361EE", "callout": "EEF0FD", "zebra": "F4F5FB", "coverSub": "C9D0F7"},
    "warm":   {"primary": "7A3E1D", "accent": "D97B29", "callout": "FDF3EA", "zebra": "FBF6F1", "coverSub": "F4D9C0"},
    "mono":   {"primary": "111111", "accent": "555555", "callout": "F2F2F2", "zebra": "F6F6F6", "coverSub": "D0D0D0"},
}

STYLES = {
    "modern":  {"font": "Calibri",         "size": 11.0, "line": 1.20, "after": 6, "h1": 16, "h2": 13, "h3": 11.5, "title": 24, "colorHeadings": True},
    "classic": {"font": "Times New Roman", "size": 12.0, "line": 1.50, "after": 6, "h1": 14, "h2": 13, "h3": 12.0, "title": 22, "colorHeadings": False},
    "compact": {"font": "Calibri",         "size": 10.5, "line": 1.15, "after": 4, "h1": 14, "h2": 12, "h3": 11.0, "title": 22, "colorHeadings": True},
}

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

TEXT_HEX = "222222"
MUTED_HEX = "666666"
LINE_HEX = "D0D7DE"


def _rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_cell_shading(cell, fill_hex: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex.lstrip("#")}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, cfg in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if cfg is None:
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:val"), "none")
            tcBorders.append(node)
        else:
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:val"), cfg.get("val", "single"))
            node.set(qn("w:sz"), str(cfg.get("sz", 4)))
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), cfg.get("color", LINE_HEX))
            tcBorders.append(node)
    tcPr.append(tcBorders)


def _set_paragraph_bottom_border(p, color_hex=LINE_HEX, sz=8, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_paragraph_left_border(p, color_hex, sz=18, space=10):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color_hex.lstrip("#"))
    pBdr.append(left)
    pPr.append(pBdr)


def _add_inline_runs(p, text, font_name, size_pt, color_hex=TEXT_HEX, bold=False, italic=False, accent_hex="2E75B6"):
    """Parse **bold**, *italic*, [label](url), dan \\n."""
    text = "" if text is None else str(text)
    token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")
    pos = 0

    def emit_plain(chunk, is_bold=bold, is_italic=italic, col=color_hex, underline=False):
        lines = chunk.split("\n")
        for i, ln in enumerate(lines):
            r = p.add_run(ln)
            r.font.name = font_name
            r.font.size = Pt(size_pt)
            r.font.bold = is_bold
            r.font.italic = is_italic
            r.font.underline = underline
            r.font.color.rgb = _rgb(col)
            if i < len(lines) - 1:
                r.add_break()

    for m in token_re.finditer(text):
        if m.start() > pos:
            emit_plain(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            emit_plain(tok[2:-2], is_bold=True)
        elif tok.startswith("*"):
            emit_plain(tok[1:-1], is_italic=True)
        else:
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if mm:
                emit_plain(mm.group(1), col=accent_hex, underline=True)
        pos = m.end()
    if pos < len(text):
        emit_plain(text[pos:])


def _add_page_number_fields(p, font_name, lang="id", include_total=True):
    lbl_page = "Page " if lang == "en" else "Halaman "
    lbl_of = " of " if lang == "en" else " dari "

    def add_txt(t):
        r = p.add_run(t)
        r.font.name = font_name
        r.font.size = Pt(9.5)
        r.font.color.rgb = _rgb(MUTED_HEX)

    def add_field(fld_name):
        r = p.add_run()
        r.font.name = font_name
        r.font.size = Pt(9.5)
        r.font.color.rgb = _rgb(MUTED_HEX)
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls("w"))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> %s </w:instrText>' % (nsdecls("w"), fld_name))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls("w"))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls("w"))
        r._r.append(fldChar1)
        r._r.append(instrText)
        r._r.append(fldChar2)
        r._r.append(fldChar3)

    add_txt(lbl_page)
    add_field("PAGE")
    if include_total:
        add_txt(lbl_of)
        add_field("NUMPAGES")


def build(spec: dict, out_path: str) -> str:
    if not isinstance(spec, dict) or not isinstance(spec.get("blocks"), list):
        raise ValueError('[word-kit] spec harus memiliki array "blocks".')

    theme = THEMES.get(spec.get("theme", "navy"), THEMES["navy"])
    st = STYLES.get(spec.get("style", "modern"), STYLES["modern"])
    lang = spec.get("lang", "id")
    pg = spec.get("page", {})
    is_letter = pg.get("size") == "Letter"
    landscape = pg.get("orientation") == "landscape"
    margin_cm = float(pg.get("margin_cm", 2.54))

    page_w_cm, page_h_cm = (21.59, 27.94) if is_letter else (21.0, 29.7)
    if landscape:
        page_w_cm, page_h_cm = page_h_cm, page_w_cm
    content_w_cm = page_w_cm - 2 * margin_cm
    content_w_in = content_w_cm / 2.54

    doc = docx.Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = st["font"]
    normal_style.font.size = Pt(st["size"])
    normal_style.font.color.rgb = _rgb(TEXT_HEX)

    def configure_section(sec):
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width = Cm(page_w_cm)
        sec.page_height = Cm(page_h_cm)
        sec.top_margin = Cm(margin_cm)
        sec.bottom_margin = Cm(margin_cm)
        sec.left_margin = Cm(margin_cm)
        sec.right_margin = Cm(margin_cm)

    configure_section(doc.sections[0])

    # Cover Page
    has_cover = bool(spec.get("cover"))
    if has_cover:
        c = spec["cover"]
        p_top = doc.add_paragraph()
        p_top.paragraph_format.space_before = Pt(110)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Cm(content_w_cm)
        _set_cell_shading(cell, theme["primary"])
        _set_cell_margins(cell, top=420, bottom=420, left=380, right=380)
        _set_cell_borders(cell)

        p0 = cell.paragraphs[0]
        if c.get("label"):
            p0.paragraph_format.space_after = Pt(8)
            r_lbl = p0.add_run(str(c["label"]).upper())
            r_lbl.font.name = st["font"]
            r_lbl.font.size = Pt(10.5)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = _rgb(theme["coverSub"])
            p_title = cell.add_paragraph()
        else:
            p_title = p0

        p_title.paragraph_format.space_after = Pt(8 if c.get("subtitle") else 0)
        r_t = p_title.add_run(str(c.get("title", "Judul Dokumen")))
        r_t.font.name = st["font"]
        r_t.font.size = Pt(st["title"])
        r_t.font.bold = True
        r_t.font.color.rgb = _rgb("FFFFFF")

        if c.get("subtitle"):
            p_sub = cell.add_paragraph()
            p_sub.paragraph_format.space_after = Pt(0)
            r_s = p_sub.add_run(str(c["subtitle"]))
            r_s.font.name = st["font"]
            r_s.font.size = Pt(14)
            r_s.font.color.rgb = _rgb(theme["coverSub"])

        p_rule = doc.add_paragraph()
        p_rule.paragraph_format.space_before = Pt(20)
        p_rule.paragraph_format.space_after = Pt(10)
        _set_paragraph_bottom_border(p_rule, theme["accent"], sz=12, space=4)

        for ln in c.get("lines", []):
            p_ln = doc.add_paragraph()
            p_ln.paragraph_format.space_after = Pt(4)
            _add_inline_runs(p_ln, ln, st["font"], 11.5, MUTED_HEX, accent_hex=theme["accent"])

        body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(body_sec)
        body_sec.header.is_linked_to_previous = False
        body_sec.footer.is_linked_to_previous = False
    else:
        body_sec = doc.sections[0]

    # Header & Footer on body section
    meta = spec.get("meta", {})
    if spec.get("header") is not False:
        hdr = body_sec.header
        hp = hdr.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(content_w_cm), WD_TAB_ALIGNMENT.RIGHT)
        _set_paragraph_bottom_border(hp, LINE_HEX, sz=4, space=4)
        h_cfg = spec.get("header") if isinstance(spec.get("header"), dict) else {}
        left_txt = h_cfg.get("left", meta.get("title", ""))
        right_txt = h_cfg.get("right", "")
        r_hl = hp.add_run(str(left_txt))
        r_hl.font.name = st["font"]
        r_hl.font.size = Pt(9.5)
        r_hl.font.color.rgb = _rgb(MUTED_HEX)
        if right_txt:
            r_hr = hp.add_run("   |   " + str(right_txt))
            r_hr.font.name = st["font"]
            r_hr.font.size = Pt(9.5)
            r_hr.font.color.rgb = _rgb(MUTED_HEX)

    if spec.get("footer") is not False:
        ftr = body_sec.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_fields(fp, st["font"], lang=lang, include_total=not has_cover)

    head_color = theme["primary"] if st["colorHeadings"] else TEXT_HEX
    h2_color = theme["accent"] if st["colorHeadings"] else TEXT_HEX

    tmp_dir = tempfile.mkdtemp(prefix="word_kit_")

    def insert_styled_image(img_path, b, idx):
        """Smart Image Payload & Placement Engine untuk Word."""
        placement = b.get("placement", "center")
        aspect = b.get("aspect_ratio")
        fit_mode = b.get("fit_mode", "contain")
        rounded = int(b.get("rounded_px", 0))
        processed_path = os.path.join(tmp_dir, f"img_{idx}.png")
        viz_kit.process_image_payload(img_path, processed_path, aspect_ratio=aspect, fit_mode=fit_mode, rounded_px=rounded)

        if placement in ("split-left", "split-right"):
            img_ratio = float(b.get("split_ratio", 0.46))
            img_w_cm = content_w_cm * img_ratio
            txt_w_cm = content_w_cm * (1.0 - img_ratio)
            tbl = doc.add_table(rows=1, cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c_left, c_right = tbl.cell(0, 0), tbl.cell(0, 1)
            _set_cell_borders(c_left)
            _set_cell_borders(c_right)
            c_left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            c_right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            if placement == "split-left":
                img_cell, txt_cell = c_left, c_right
                c_left.width, c_right.width = Cm(img_w_cm), Cm(txt_w_cm)
                _set_cell_margins(c_left, top=60, bottom=60, left=0, right=140)
                _set_cell_margins(c_right, top=60, bottom=60, left=140, right=0)
            else:
                txt_cell, img_cell = c_left, c_right
                c_left.width, c_right.width = Cm(txt_w_cm), Cm(img_w_cm)
                _set_cell_margins(c_left, top=60, bottom=60, left=0, right=140)
                _set_cell_margins(c_right, top=60, bottom=60, left=140, right=0)

            ip = img_cell.paragraphs[0]
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.space_after = Pt(4)
            ip.add_run().add_picture(processed_path, width=Cm(img_w_cm - 0.5))
            if b.get("caption"):
                cp = img_cell.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(0)
                _add_inline_runs(cp, b["caption"], st["font"], 9.5, MUTED_HEX, italic=True, accent_hex=theme["accent"])

            tp = txt_cell.paragraphs[0]
            if b.get("side_title"):
                tp.paragraph_format.space_after = Pt(6)
                _add_inline_runs(tp, b["side_title"], st["font"], st["h2"], h2_color, bold=True, accent_hex=theme["accent"])
                tp = txt_cell.add_paragraph()
            if b.get("side_text"):
                tp.paragraph_format.space_after = Pt(6)
                tp.paragraph_format.line_spacing = st["line"]
                _add_inline_runs(tp, b["side_text"], st["font"], st["size"], TEXT_HEX, accent_hex=theme["accent"])
            for bul in b.get("side_bullets", []):
                bp = txt_cell.add_paragraph()
                bp.paragraph_format.left_indent = Inches(0.2)
                bp.paragraph_format.space_after = Pt(4)
                _add_inline_runs(bp, "• " + str(bul), st["font"], st["size"], TEXT_HEX, accent_hex=theme["accent"])
            return

        # Standard placements: full-banner, center, left, right
        if b.get("width_cm"):
            w_cm = min(float(b["width_cm"]), content_w_cm)
        elif b.get("width"):
            w_cm = min(float(b["width"]) / 37.8, content_w_cm)
        elif placement == "full-banner":
            w_cm = content_w_cm
        else:
            w_cm = min(content_w_cm * 0.82, 13.5)

        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3 if b.get("caption") else 8)
        if b.get("caption"):
            p_img.paragraph_format.keep_with_next = True
        p_img.alignment = ALIGN_MAP.get(placement if placement in ("left", "right", "center") else "center", WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(processed_path, width=Cm(w_cm))

        if b.get("caption"):
            p_cap = doc.add_paragraph()
            p_cap.alignment = p_img.alignment
            p_cap.paragraph_format.space_after = Pt(8)
            _add_inline_runs(p_cap, b["caption"], st["font"], 9.5, MUTED_HEX, italic=True, accent_hex=theme["accent"])

    # Render blocks
    for idx, b in enumerate(spec["blocks"]):
        t = b.get("t")
        if t == "title":
            p = doc.add_paragraph()
            p.alignment = ALIGN_MAP.get(b.get("align", "left"), WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(4 if b.get("subtitle") else 8)
            p.paragraph_format.keep_with_next = True
            _add_inline_runs(p, b.get("text", ""), st["font"], st["title"], head_color, bold=True, accent_hex=theme["accent"])
            if b.get("subtitle"):
                ps = doc.add_paragraph()
                ps.alignment = p.alignment
                ps.paragraph_format.space_after = Pt(8)
                ps.paragraph_format.keep_with_next = True
                _add_inline_runs(ps, b["subtitle"], st["font"], 13, MUTED_HEX, accent_hex=theme["accent"])
            pr = doc.add_paragraph()
            pr.paragraph_format.space_after = Pt(12)
            _set_paragraph_bottom_border(pr, theme["accent"], sz=12, space=4)

        elif t in ("h1", "h2", "h3"):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            sz = st[t]
            col = head_color if t == "h1" else (h2_color if t == "h2" else TEXT_HEX)
            p.paragraph_format.space_before = Pt(14 if t == "h1" else (11 if t == "h2" else 8))
            p.paragraph_format.space_after = Pt(5 if t == "h1" else 4)
            _add_inline_runs(p, b.get("text", ""), st["font"], sz, col, bold=True, accent_hex=theme["accent"])

        elif t == "p":
            p = doc.add_paragraph()
            p.alignment = ALIGN_MAP.get(b.get("align", "left"), WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.line_spacing = st["line"]
            p.paragraph_format.space_after = Pt(b.get("after", st["after"]))
            if b.get("keep"):
                p.paragraph_format.keep_together = True
            sz = float(b["size"]) / 2.0 if b.get("size") and float(b["size"]) >= 16 else float(b.get("size_pt", st["size"]))
            sz = max(sz, 10.0)
            col = b.get("color", TEXT_HEX)
            _add_inline_runs(p, b.get("text", ""), st["font"], sz, col, bold=bool(b.get("bold")), italic=bool(b.get("italic")), accent_hex=theme["accent"])

        elif t == "bullets":
            items = b.get("items", [])
            for it_i, it in enumerate(items):
                txt = it if isinstance(it, str) else it.get("text", "")
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6 if it_i == len(items) - 1 and not (isinstance(it, dict) and it.get("sub")) else 3)
                _add_inline_runs(p, txt, st["font"], st["size"], TEXT_HEX, accent_hex=theme["accent"])
                if isinstance(it, dict) and isinstance(it.get("sub"), list):
                    for s_i, sub_txt in enumerate(it["sub"]):
                        ps = doc.add_paragraph()
                        ps.paragraph_format.left_indent = Inches(0.5)
                        ps.paragraph_format.line_spacing = 1.15
                        ps.paragraph_format.space_after = Pt(6 if it_i == len(items) - 1 and s_i == len(it["sub"]) - 1 else 2)
                        _add_inline_runs(ps, "– " + str(sub_txt), st["font"], st["size"], TEXT_HEX, accent_hex=theme["accent"])

        elif t == "numbers":
            items = b.get("items", [])
            for num_i, it in enumerate(items, start=1):
                txt = it if isinstance(it, str) else it.get("text", "")
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.20)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6 if num_i == len(items) else 3)
                _add_inline_runs(p, f"**{num_i}.** {txt}", st["font"], st["size"], TEXT_HEX, accent_hex=theme["accent"])

        elif t == "table":
            headers = b.get("header", [])
            rows = b.get("rows", [])
            n_cols = len(headers)
            if n_cols == 0:
                continue
            widths_raw = b.get("widths")
            if not isinstance(widths_raw, list) or len(widths_raw) != n_cols:
                widths_raw = [1.0] * n_cols
            w_sum = sum(float(x) for x in widths_raw)
            col_cms = [content_w_cm * (float(x) / w_sum) for x in widths_raw]
            aligns = b.get("align", [])
            zebra = b.get("zebra", True)

            tbl = doc.add_table(rows=1 + len(rows), cols=n_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            border_cfg = {"val": "single", "sz": 4, "color": LINE_HEX}

            # Header row
            hdr_row = tbl.rows[0]
            hdr_trPr = hdr_row._tr.get_or_add_trPr()
            hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for c_i, h_text in enumerate(headers):
                cell = hdr_row.cells[c_i]
                cell.width = Cm(col_cms[c_i])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_shading(cell, theme["primary"])
                _set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
                _set_cell_borders(cell, border_cfg, border_cfg, border_cfg, border_cfg)
                cp = cell.paragraphs[0]
                cp.alignment = ALIGN_MAP.get(aligns[c_i] if c_i < len(aligns) else "left", WD_ALIGN_PARAGRAPH.LEFT)
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.keep_with_next = True
                _add_inline_runs(cp, h_text, st["font"], 10.5, "FFFFFF", bold=True, accent_hex=theme["accent"])

            # Data rows (strict >= 10pt floor)
            for r_i, r_vals in enumerate(rows):
                row = tbl.rows[r_i + 1]
                trPr = row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                is_bold_row = bool(b.get("boldLastRow") and r_i == len(rows) - 1)
                for c_i in range(n_cols):
                    val = r_vals[c_i] if c_i < len(r_vals) else ""
                    cell = row.cells[c_i]
                    cell.width = Cm(col_cms[c_i])
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    if zebra and r_i % 2 == 1:
                        _set_cell_shading(cell, theme["zebra"])
                    _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    _set_cell_borders(cell, border_cfg, border_cfg, border_cfg, border_cfg)
                    cp = cell.paragraphs[0]
                    cp.alignment = ALIGN_MAP.get(aligns[c_i] if c_i < len(aligns) else "left", WD_ALIGN_PARAGRAPH.LEFT)
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    if len(rows) <= 5 and r_i < len(rows) - 1:
                        cp.paragraph_format.keep_with_next = True
                    _add_inline_runs(cp, val, st["font"], 10.0, TEXT_HEX, bold=is_bold_row, accent_hex=theme["accent"])

            if b.get("caption"):
                p_cap = doc.add_paragraph()
                p_cap.paragraph_format.space_before = Pt(4)
                p_cap.paragraph_format.space_after = Pt(8)
                _add_inline_runs(p_cap, b["caption"], st["font"], 9.5, MUTED_HEX, italic=True, accent_hex=theme["accent"])
            else:
                doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif t == "kv":
            rows = b.get("rows", [])
            w_raw = b.get("widths", [28, 72])
            w_sum = sum(float(x) for x in w_raw)
            w0, w1 = content_w_cm * (w_raw[0] / w_sum), content_w_cm * (w_raw[1] / w_sum)
            tbl = doc.add_table(rows=len(rows), cols=2)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            b_bot = {"val": "single", "sz": 4, "color": LINE_HEX} if b.get("lines") else None
            for r_i, pair in enumerate(rows):
                row = tbl.rows[r_i]
                row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                c0, c1 = row.cells[0], row.cells[1]
                c0.width, c1.width = Cm(w0), Cm(w1)
                _set_cell_margins(c0, top=50, bottom=50, left=60, right=60)
                _set_cell_margins(c1, top=50, bottom=50, left=60, right=60)
                _set_cell_borders(c0, bottom=b_bot)
                _set_cell_borders(c1, bottom=b_bot)
                p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
                p0.paragraph_format.space_after = Pt(2)
                p1.paragraph_format.space_after = Pt(2)
                if len(rows) <= 6 and r_i < len(rows) - 1:
                    p0.paragraph_format.keep_with_next = True
                    p1.paragraph_format.keep_with_next = True
                _add_inline_runs(p0, pair[0] if len(pair) > 0 else "", st["font"], 10.5, b.get("keyColor", theme["primary"]), bold=True, accent_hex=theme["accent"])
                _add_inline_runs(p1, pair[1] if len(pair) > 1 else "", st["font"], 10.5, TEXT_HEX, accent_hex=theme["accent"])

        elif t == "kpi_row":
            cards = b.get("cards", [])
            if not cards:
                continue
            n_c = len(cards)
            cw = content_w_cm / n_c
            tbl = doc.add_table(rows=1, cols=n_c)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            b_card = {"val": "single", "sz": 6, "color": LINE_HEX}
            b_top = {"val": "single", "sz": 18, "color": theme["accent"]}
            for ci, card in enumerate(cards):
                cell = tbl.cell(0, ci)
                cell.width = Cm(cw)
                _set_cell_shading(cell, theme["callout"])
                _set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                _set_cell_borders(cell, top=b_top, bottom=b_card, left=b_card, right=b_card)
                p_lbl = cell.paragraphs[0]
                p_lbl.paragraph_format.space_after = Pt(2)
                _add_inline_runs(p_lbl, str(card.get("label", "")).upper(), st["font"], 10.0, MUTED_HEX, bold=True)
                p_val = cell.add_paragraph()
                p_val.paragraph_format.space_after = Pt(2)
                _add_inline_runs(p_val, str(card.get("value", "")), st["font"], 16.0, theme["primary"], bold=True)
                if card.get("sub"):
                    p_sub = cell.add_paragraph()
                    p_sub.paragraph_format.space_after = Pt(0)
                    _add_inline_runs(p_sub, str(card["sub"]), st["font"], 10.0, theme["accent"], bold=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif t == "callout":
            warn = b.get("kind") == "warning"
            bar_col = "C2410C" if warn else theme["accent"]
            fill_col = "FEF3E8" if warn else theme["callout"]
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            cell = tbl.cell(0, 0)
            cell.width = Cm(content_w_cm)
            _set_cell_shading(cell, fill_col)
            _set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
            _set_cell_borders(cell, left={"val": "single", "sz": 24, "color": bar_col})
            p0 = cell.paragraphs[0]
            if b.get("title"):
                p0.paragraph_format.space_after = Pt(3)
                _add_inline_runs(p0, b["title"], st["font"], 11.0, bar_col, bold=True, accent_hex=theme["accent"])
                p_body = cell.add_paragraph()
            else:
                p_body = p0
            p_body.paragraph_format.space_after = Pt(0)
            p_body.paragraph_format.line_spacing = 1.15
            _add_inline_runs(p_body, b.get("text", ""), st["font"], 10.5, TEXT_HEX, accent_hex=theme["accent"])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif t == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            _set_paragraph_left_border(p, theme["accent"], sz=18, space=10)
            _add_inline_runs(p, b.get("text", ""), st["font"], st["size"], "444444", italic=True, accent_hex=theme["accent"])
            if b.get("source"):
                _add_inline_runs(p, "  — " + str(b["source"]), st["font"], 10.0, MUTED_HEX, accent_hex=theme["accent"])

        elif t == "code":
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Cm(content_w_cm)
            _set_cell_shading(cell, "F5F5F5")
            _set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            _set_cell_borders(cell, left={"val": "single", "sz": 12, "color": LINE_HEX})
            lines = str(b.get("text", "")).split("\n")
            for li, ln in enumerate(lines):
                cp = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                cp.paragraph_format.space_after = Pt(1)
                r = cp.add_run(ln if ln else " ")
                r.font.name = "Courier New"
                r.font.size = Pt(10.0)
                r.font.color.rgb = _rgb(TEXT_HEX)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif t == "image":
            img_path = b.get("path")
            if not img_path or not os.path.exists(img_path):
                # Auto-generate diagram payload if diagram_spec provided
                if isinstance(b.get("diagram_spec"), dict):
                    img_path = os.path.join(tmp_dir, f"auto_diag_{idx}.png")
                    d_spec = {"theme": spec.get("theme", "navy"), "width_px": 1100, "height_px": 950, **b["diagram_spec"]}
                    viz_kit.generate_diagram_payload(d_spec, img_path)
                else:
                    raise FileNotFoundError(f"[word-kit] Blok #{idx+1} (image): file tidak ditemukan: {img_path}")
            insert_styled_image(img_path, b, idx)

        elif t in ("curve", "chart"):
            curve_png = os.path.join(tmp_dir, f"curve_{idx}.png")
            c_spec = {"theme": spec.get("theme", "navy"), **b}
            viz_kit.render_curve(c_spec, curve_png)
            insert_styled_image(curve_png, b, idx)

        elif t == "gallery":
            items = b.get("items", [])
            cols = int(b.get("cols", min(len(items), 3) or 2))
            aspect = b.get("aspect_ratio", "4:3")
            n_rows = (len(items) + cols - 1) // cols
            tbl = doc.add_table(rows=n_rows, cols=cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cw = content_w_cm / cols
            for ii, item in enumerate(items):
                ri, ci = divmod(ii, cols)
                cell = tbl.cell(ri, ci)
                cell.width = Cm(cw)
                _set_cell_borders(cell)
                _set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
                p_in = item.get("path")
                out_p = os.path.join(tmp_dir, f"gal_{idx}_{ii}.png")
                viz_kit.process_image_payload(p_in, out_p, target_w_px=900, aspect_ratio=aspect, fit_mode="cover", rounded_px=int(b.get("rounded_px", 12)))
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(3)
                cp.add_run().add_picture(out_p, width=Cm(cw - 0.5))
                if item.get("caption"):
                    cap_p = cell.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(2)
                    _add_inline_runs(cap_p, item["caption"], st["font"], 9.5, MUTED_HEX, italic=True)

        elif t == "pagebreak":
            doc.add_page_break()

        elif t == "spacer":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(float(b.get("size", 120)) / 20.0)

        elif t == "rule":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _set_paragraph_bottom_border(p, b.get("color", LINE_HEX), sz=int(b.get("size", 8)), space=4)

        elif t == "sectionhead":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(5)
            _set_paragraph_bottom_border(p, theme["accent"], sz=8, space=3)
            _add_inline_runs(p, str(b.get("text", "")).upper(), st["font"], 11.0, theme["primary"], bold=True)

        elif t == "lr":
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(content_w_cm), WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.keep_with_next = bool(b.get("sub"))
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1 if b.get("sub") else 3)
            _add_inline_runs(p, b.get("left", ""), st["font"], st["size"], TEXT_HEX, bold=b.get("bold", True), accent_hex=theme["accent"])
            r_tab = p.add_run("\t")
            r_tab.font.name = st["font"]
            _add_inline_runs(p, b.get("right", ""), st["font"], st["size"], MUTED_HEX, accent_hex=theme["accent"])
            if b.get("sub"):
                ps = doc.add_paragraph()
                ps.paragraph_format.space_after = Pt(4)
                _add_inline_runs(ps, b["sub"], st["font"], 10.5, MUTED_HEX, italic=True, accent_hex=theme["accent"])

        elif t == "letterhead":
            p_name = doc.add_paragraph()
            p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_name.paragraph_format.space_after = Pt(2)
            _add_inline_runs(p_name, str(b.get("name", "")).upper(), st["font"], 16.0, theme["primary"], bold=True)
            lines = b.get("lines", [])
            for li, ln in enumerate(lines):
                pl = doc.add_paragraph()
                pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pl.paragraph_format.space_after = Pt(8 if li == len(lines) - 1 else 1)
                if li == len(lines) - 1:
                    _set_paragraph_bottom_border(pl, theme["primary"], sz=12, space=6)
                _add_inline_runs(pl, ln, st["font"], 10.0, MUTED_HEX, accent_hex=theme["accent"])

        elif t == "signature":
            if b.get("closing"):
                pc = doc.add_paragraph()
                pc.paragraph_format.space_before = Pt(12)
                pc.paragraph_format.space_after = Pt(36)
                pc.paragraph_format.keep_with_next = True
                _add_inline_runs(pc, b["closing"], st["font"], st["size"], TEXT_HEX)
            pn = doc.add_paragraph()
            pn.paragraph_format.space_after = Pt(2)
            pn.paragraph_format.keep_with_next = bool(b.get("title"))
            _add_inline_runs(pn, b.get("name", ""), st["font"], st["size"], TEXT_HEX, bold=True)
            if b.get("title"):
                pt = doc.add_paragraph()
                pt.paragraph_format.space_after = Pt(4)
                _add_inline_runs(pt, b["title"], st["font"], st["size"], MUTED_HEX)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Pemakaian: python3 word-kit.py spec.json output.docx", file=sys.stderr)
        sys.exit(1)
    spec_p, out_p = sys.argv[1], sys.argv[2]
    with open(spec_p, "r", encoding="utf-8") as f:
        spec_obj = json.load(f)
    res = build(spec_obj, out_p)
    print(f"OK -> {res}")
